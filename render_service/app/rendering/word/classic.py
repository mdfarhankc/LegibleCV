from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph as DocxParagraph

from app.rendering.richtext import Paragraph as RichParagraph
from app.rendering.styles import ClassicStyle
from app.rendering.view import (
    SEPARATOR,
    Block,
    BulletsBlock,
    EntriesBlock,
    EntryView,
    InlineBlock,
    LinkText,
    ResumeView,
    SkillsBlock,
    TextBlock,
)
from app.rendering.word.helpers import (
    add_hyperlink,
    add_rich_text,
    add_right_tab,
    add_text,
    set_bottom_border,
    set_page,
    set_spacing,
)


class ClassicBuilder:
    def __init__(self, document: Document, view: ResumeView, style: ClassicStyle) -> None:
        self._doc = document
        self._view = view
        self._style = style
        self._tab_in = style.content_width_in(view.page)

    def build(self) -> None:
        self._configure()
        self._header()
        for section in self._view.sections:
            self._heading(section.title)
            self._block(section.block)

    def _configure(self) -> None:
        style, view = self._style, self._view
        set_page(
            self._doc,
            width_in=view.page.width_in,
            height_in=view.page.height_in,
            margin_x_in=style.margin_x_in,
            margin_y_in=style.margin_y_in,
        )
        normal = self._doc.styles["Normal"]
        normal.font.name = style.docx_font
        normal.font.size = Pt(style.body_pt)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = style.line_height

        props = self._doc.core_properties
        props.title = view.document_title
        props.author = view.full_name
        props.last_modified_by = view.full_name
        props.comments = ""
        props.keywords = ""
        props.subject = ""

    def _header(self) -> None:
        style = self._style
        name = self._doc.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(name, self._view.full_name, size_pt=style.name_pt, bold=True)
        set_spacing(name, after_pt=style.name_after_pt)

        contact = self._doc.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for index, item in enumerate(self._view.contact):
            if index:
                add_text(contact, SEPARATOR, size_pt=style.contact_pt)
            self._link(contact, item, size_pt=style.contact_pt)
        set_spacing(contact, after_pt=style.contact_after_pt)

    def _heading(self, title: str) -> None:
        style = self._style
        heading = self._doc.add_paragraph()
        add_text(heading, title, size_pt=style.heading_pt, bold=True)
        set_spacing(
            heading,
            before_pt=style.heading_before_pt,
            after_pt=style.heading_after_pt,
            keep_with_next=True,
        )
        set_bottom_border(heading, width_pt=style.heading_border_pt, color=style.text_color)

    def _block(self, block: Block) -> None:
        match block:
            case TextBlock(paragraphs=paragraphs):
                for paragraph in paragraphs:
                    self._rich_paragraph(paragraph)
            case EntriesBlock(entries=entries):
                for index, entry in enumerate(entries):
                    self._entry(entry, first=index == 0)
            case SkillsBlock(groups=groups):
                for label, items in groups:
                    self._skill(label, items)
            case BulletsBlock(items=items):
                for item in items:
                    self._bullet(item)
            case InlineBlock(items=items):
                self._inline(items)

    def _entry(self, entry: EntryView, *, first: bool) -> None:
        style = self._style
        title = self._doc.add_paragraph()
        add_right_tab(title, self._tab_in)
        add_text(title, entry.title, size_pt=style.entry_title_pt, bold=True)
        if entry.title_right:
            add_text(title, f"\t{entry.title_right}", size_pt=style.body_pt)
        set_spacing(
            title,
            before_pt=0 if first else style.entry_before_pt,
            after_pt=style.line_after_pt,
            keep_with_next=True,
        )

        if entry.has_subtitle_line:
            subtitle = self._doc.add_paragraph()
            add_right_tab(subtitle, self._tab_in)
            if entry.subtitle:
                add_text(subtitle, entry.subtitle, size_pt=style.body_pt, italic=True)
            if entry.subtitle_right:
                add_text(subtitle, f"\t{entry.subtitle_right}", size_pt=style.body_pt)
            set_spacing(subtitle, after_pt=style.line_after_pt, keep_with_next=True)

        if entry.meta:
            meta = self._doc.add_paragraph()
            add_text(
                meta,
                entry.meta,
                size_pt=style.meta_pt,
                italic=True,
                color=style.muted(self._view.ats_mode),
            )
            set_spacing(meta, after_pt=style.line_after_pt, keep_with_next=True)

        if entry.link:
            link = self._doc.add_paragraph()
            self._link(link, entry.link, size_pt=style.body_pt)
            set_spacing(link, after_pt=style.line_after_pt, keep_with_next=bool(entry.bullets))

        for paragraph in entry.description:
            self._rich_paragraph(paragraph)
        for bullet in entry.bullets:
            self._bullet(bullet)

    def _rich_paragraph(self, runs: RichParagraph) -> None:
        paragraph = self._doc.add_paragraph()
        add_rich_text(paragraph, runs, size_pt=self._style.body_pt)
        set_spacing(paragraph, after_pt=self._style.paragraph_after_pt)

    def _bullet(self, runs: RichParagraph) -> None:
        style = self._style
        bullet = self._doc.add_paragraph(style="List Bullet")
        add_rich_text(bullet, runs, size_pt=style.body_pt)
        set_spacing(bullet, before_pt=style.bullet_gap_pt, after_pt=style.bullet_gap_pt)

    def _skill(self, label: str, items: str) -> None:
        style = self._style
        skill = self._doc.add_paragraph()
        add_text(skill, f"{label}: ", size_pt=style.body_pt, bold=True)
        add_text(skill, items, size_pt=style.body_pt)
        set_spacing(skill, after_pt=style.line_after_pt)

    def _inline(self, items: list[str]) -> None:
        style = self._style
        paragraph = self._doc.add_paragraph()
        add_text(paragraph, SEPARATOR.join(items), size_pt=style.body_pt)
        set_spacing(paragraph, after_pt=style.line_after_pt)

    def _link(self, paragraph: DocxParagraph, item: LinkText, *, size_pt: float) -> None:
        if item.url is None:
            add_text(paragraph, item.text, size_pt=size_pt)
            return
        add_hyperlink(
            paragraph,
            item.text,
            item.url,
            size_pt=size_pt,
            color=self._style.link(self._view.ats_mode),
            underline=not self._view.ats_mode,
        )
