from io import BytesIO
from typing import Any

import docx
import pytest
from docx.shared import Inches

from app.rendering.word import render_docx
from app.schemas.resume import Resume
from tests.helpers import FIXTURE, ORDERED_MARKERS, assert_in_order, docx_text


@pytest.fixture(scope="module")
def sample_docx() -> bytes:
    return render_docx(Resume.model_validate_json(FIXTURE.read_text(encoding="utf-8")))


def _open(content: bytes) -> docx.document.Document:
    return docx.Document(BytesIO(content))


def test_text_is_ordered(sample_docx: bytes) -> None:
    assert_in_order(docx_text(sample_docx), ORDERED_MARKERS)


def test_bullets_use_list_style_and_bold_runs(sample_docx: bytes) -> None:
    document = _open(sample_docx)
    bullets = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 25
    bold = [run.text for run in bullets[1].runs if run.bold]
    assert bold == ["FastAPI", "Django"]


def test_right_aligned_dates_use_tab(sample_docx: bytes) -> None:
    document = _open(sample_docx)
    job = next(p for p in document.paragraphs if p.text.startswith("NORTHWIND"))
    assert job.text == "NORTHWIND SYSTEMS LLC\tSpringfield, IL | Remote"


def test_hyperlinks(sample_docx: bytes) -> None:
    document = _open(sample_docx)
    urls = {link.address for paragraph in document.paragraphs for link in paragraph.hyperlinks}
    assert "https://github.com/alex-morgan-example" in urls
    assert "mailto:alex.morgan@example.com" in urls


def test_headings_have_bottom_border(sample_docx: bytes) -> None:
    document = _open(sample_docx)
    heading = next(p for p in document.paragraphs if p.text == "WORK EXPERIENCE")
    assert heading._p.pPr.find(docx.oxml.ns.qn("w:pBdr")) is not None
    assert heading.paragraph_format.keep_with_next is True


def test_metadata(sample_docx: bytes) -> None:
    props = _open(sample_docx).core_properties
    assert props.title == "ALEX MORGAN - Resume"
    assert props.author == "ALEX MORGAN"
    assert props.last_modified_by == "ALEX MORGAN"


@pytest.mark.parametrize(("page_size", "width_in"), [("A4", 8.268), ("Letter", 8.5)])
def test_page_size(resume_data: dict[str, Any], page_size: str, width_in: float) -> None:
    resume_data["options"] = {"page_size": page_size}
    section = _open(render_docx(Resume.model_validate(resume_data))).sections[0]
    assert section.page_width.inches == pytest.approx(width_in, abs=0.001)
    assert section.left_margin == Inches(0.6)


def test_user_input_is_plain_text(resume_data: dict[str, Any]) -> None:
    resume_data["full_name"] = "<script>alert(1)</script>"
    assert "<script>alert(1)</script>" in docx_text(render_docx(Resume.model_validate(resume_data)))


def test_ats_mode_prints_full_urls(resume_data: dict[str, Any]) -> None:
    resume_data["options"] = {"ats_mode": True}
    text = docx_text(render_docx(Resume.model_validate(resume_data)))
    assert "https://www.linkedin.com/in/alex-morgan-example" in text
